"""Controlli offline; dati artificiali, nessuna misura reale."""

import argparse
import asyncio
from collections import deque
from datetime import datetime, timezone
import io
import json
from pathlib import Path
import struct
import sys
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import patch
import zipfile

import garmin_api
import logger


def fixture_fit():
    # FIT minimale con record HR originali: il decoder reale verifica anche il CRC.
    def crc(data):
        value = 0
        for byte in data:
            value ^= byte
            for _ in range(8):
                value = (value >> 1) ^ (0xA001 if value & 1 else 0)
        return value

    epoch = int(datetime(2026, 10, 13, 12, tzinfo=timezone.utc).timestamp())
    fit_epoch = epoch - 631065600
    definition = bytes([0x40, 0, 0]) + struct.pack("<H", 20) + bytes([2, 253, 4, 0x86, 3, 1, 2])
    records = b"".join(bytes([0]) + struct.pack("<IB", fit_epoch + i, hr)
                       for i, hr in enumerate((70, 71, 255, 0, 74)))
    data = definition + records
    head = struct.pack("<BBHI4s", 14, 0x20, 21800, len(data), b".FIT")
    head += struct.pack("<H", crc(head))
    payload = head + data
    return payload + struct.pack("<H", crc(payload)), epoch


class LoggerTests(unittest.TestCase):
    def test_reference_formats_and_invalid_values(self):
        self.assertEqual(logger.parse_microbit("BPM 73")["bpm"], 73)
        value = logger.parse_microbit("512 498 74 extra")
        self.assertEqual((value["signal"], value["filtered"], value["bpm"]), (512, 498, 74))
        self.assertEqual(value["raw"], "512 498 74 extra")
        for text in ("BPM 0", "BPM NaN", "BPM 75.5", "BPM", "signal filtered 75", "BPM 73junk"):
            self.assertIsNone(logger.parse_microbit(text))

    def test_uart_fragmentation_with_line_framing(self):
        framer = logger.Framer("line")
        self.assertEqual(framer.feed(b"BPM 7"), [])
        self.assertEqual(framer.feed(b"3\r\n512 498"), ["BPM 73"])
        self.assertEqual(framer.feed(b" 74\nBPM 75\n"), ["512 498 74", "BPM 75"])

    def test_no_stale_microbit_and_tie_is_previous(self):
        samples = deque([logger.Sample(9, 1009, {}), logger.Sample(11, 1011, {})])
        self.assertEqual(logger.nearest_microbit(samples, 10, 1).mono, 9)
        self.assertIsNone(logger.nearest_microbit(samples, 20, 1))

    def test_timestamp_join_preserves_missing_and_original_times(self):
        session = {"rows": [
            {"n": 1, "target_epoch": 1000, "micro_epoch": 1000.2},
            {"n": 2, "target_epoch": 1010, "micro_epoch": 1010.1},
            {"n": 3, "target_epoch": 1020},
            {"n": 4, "target_epoch": 1030, "micro_epoch": 1030},
        ], "metadata": {}}
        logger.match_garmin(session, [(995, 70), (996, 71), (1004, 72), (5000, 99)], 1, 5)
        first, second, third, fourth = session["rows"]
        self.assertEqual(first["garmin_original_epoch"], 995)
        self.assertEqual(first["garmin_epoch"], 1000)
        self.assertEqual(first["garmin_bpm"], 70)
        self.assertEqual(first["delta_seconds"], -0.2)
        self.assertNotIn("garmin_bpm", second)
        self.assertEqual(third["status"], "MICROBIT_MANCANTE")
        self.assertNotIn("garmin_bpm", fourth)
        # Ripetere l'importazione non deve lasciare valori della vecchia sorgente.
        logger.match_garmin(session, [], 1)
        self.assertNotIn("garmin_bpm", first)

    def test_grid_and_actual_sample_must_both_match(self):
        session = {"rows": [{"n": 1, "target_epoch": 1000, "micro_epoch": 1000.9}], "metadata": {}}
        logger.match_garmin(session, [(1001.5, 75)], 1)
        self.assertNotIn("garmin_bpm", session["rows"][0])

    def test_exact_matching_and_no_reuse(self):
        session = {"rows": [{"n": i, "target_epoch": 1000, "micro_epoch": 1000} for i in (1, 2)], "metadata": {}}
        logger.match_garmin(session, [(1000, 75)], 0)
        self.assertEqual(session["rows"][0]["garmin_bpm"], 75)
        self.assertNotIn("garmin_bpm", session["rows"][1])

    def test_excel_values_dates_and_text_are_readable(self):
        from openpyxl import load_workbook  # Solo lettura per verificare il file.
        session = {"rows": [{"n": 1, "target_epoch": 1791892800.125,
                              "micro_epoch": 1791892800.25, "micro_bpm": 73,
                              "status": "GARMIN_DA_SCARICARE", "raw": "=TEST\x00<&>"}],
                   "metadata": {"origine": "dati artificiali di test"}}
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "test.xlsx"
            logger.write_xlsx(output, session)
            workbook = load_workbook(output, read_only=True, data_only=False)
            self.assertEqual(workbook.sheetnames, ["Misure", "Impostazioni"])
            cells = list(workbook["Misure"].iter_rows())[1]
            self.assertEqual(cells[3].value, 73)
            self.assertIsNone(cells[8].value)
            self.assertEqual(cells[11].data_type, "s")
            self.assertEqual(cells[11].value, "=TEST<&>")
            expected = datetime.fromtimestamp(1791892800.125, timezone.utc).replace(tzinfo=None)
            self.assertLess(abs((cells[1].value - expected).total_seconds()), 0.002)
            workbook.close()

    def test_locked_excel_keeps_session_and_reports_failure(self):
        with tempfile.TemporaryDirectory() as directory:
            target = Path(directory) / "session.json"
            with patch.object(logger, "write_xlsx", side_effect=PermissionError), patch("sys.stderr", new=io.StringIO()):
                self.assertFalse(logger.checkpoint(Path(directory) / "test.xlsx", target, {"rows": []}))
            self.assertEqual(json.loads(target.read_text())["rows"], [])

    def test_real_fit_decoder_original_records_and_crc(self):
        data, epoch = fixture_fit()
        self.assertEqual(garmin_api.read_fit_samples(data), [(epoch, 70), (epoch + 1, 71), (epoch + 4, 74)])
        broken = bytearray(data)
        broken[-3] ^= 1
        with self.assertRaises(ValueError):
            garmin_api.read_fit_samples(bytes(broken))

    def test_actual_client_methods_exist_without_login(self):
        from garminconnect import Garmin
        import inspect
        self.assertIn("tokenstore", inspect.signature(Garmin.login).parameters)
        self.assertIn("dl_fmt", inspect.signature(Garmin.download_activity).parameters)
        self.assertIn("prompt_mfa", inspect.signature(Garmin).parameters)
        self.assertIsNotNone(Garmin.ActivityDownloadFormat.ORIGINAL)

    def test_recording_grid_and_dropped_notification_with_fake_device(self):
        clock = SimpleNamespace(mono=0.0, epoch=1791892800.0, notify=None, next_event=10)
        char = SimpleNamespace(uuid=logger.UART_NOTIFY, properties=["indicate"])
        service = SimpleNamespace(characteristics=[char])

        class FakeClient:
            def __init__(self, *_args, **_kwargs):
                self.services = SimpleNamespace(get_service=lambda _uuid: service)

            async def __aenter__(self):
                return self

            async def __aexit__(self, *_args):
                pass

            async def start_notify(self, _char, callback):
                clock.notify = callback
                callback(char, bytearray(b"BPM 70\n"))

        async def fake_discover(_args):
            return SimpleNamespace(address="TEST-MICROBIT")

        async def fake_sleep(seconds):
            end = clock.mono + seconds
            while clock.next_event <= end:
                clock.mono = clock.next_event
                if clock.next_event != 20:  # Perdita del campione a t0 + 20s.
                    clock.notify(char, bytearray(f"BPM {70 + clock.next_event // 10}\n".encode()))
                clock.next_event += 10
            clock.mono = end

        with tempfile.TemporaryDirectory() as directory:
            args = argparse.Namespace(output=Path(directory) / "test.xlsx", framing="line", warmup=0,
                                      tolerance=1, samples=4, pair=False, characteristic=None)
            fake_bleak = SimpleNamespace(BleakClient=FakeClient)
            with patch.dict(sys.modules, {"bleak": fake_bleak}), \
                 patch.object(logger, "discover", new=fake_discover), \
                 patch.object(logger.time, "monotonic", side_effect=lambda: clock.mono), \
                 patch.object(logger.time, "time", side_effect=lambda: clock.epoch + clock.mono), \
                 patch.object(logger.asyncio, "sleep", new=fake_sleep), \
                 patch("sys.stdout", new=io.StringIO()):
                asyncio.run(logger.record(args))
            session = logger.load_session(args.output.with_suffix(".session.json"))
            self.assertEqual([r["target_epoch"] - clock.epoch for r in session["rows"]], [0, 10, 20, 30])
            self.assertEqual([r.get("micro_bpm") for r in session["rows"]], [70, 71, None, 73])
            self.assertEqual(session["rows"][2]["status"], "MICROBIT_MANCANTE")
            self.assertEqual(session["metadata"]["stato_registrazione"], "completata")


if __name__ == "__main__":
    unittest.main()
