"""Analisi Lab 6: artificiali a 30 e 150 bpm separati, Pietro e Manfred.

Esecuzione: python analisi_lab6.py
Dipendenze: requirements_analisi.txt. Il report HTML include i grafici e funziona offline.
"""

from __future__ import annotations

import argparse
import base64
from collections import Counter
import csv
from dataclasses import dataclass
from datetime import datetime, timezone
import hashlib
from html import escape
import importlib.metadata
import json
import math
from pathlib import Path
import sys

# Librerie locali installate per questa analisi, senza cambiare il vecchio venv.
ROOT = Path(__file__).resolve().parent
LOCAL_PACKAGES = ROOT / ".analysis_packages"
if LOCAL_PACKAGES.is_dir():
    sys.path.insert(0, str(LOCAL_PACKAGES))

try:
    import numpy as np
    from openpyxl import load_workbook
    from scipy import stats
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
except ImportError as exc:
    raise SystemExit(
        "Dipendenza mancante. Installa requirements_analisi.txt con il Python "
        "che esegue questo file. Dettaglio: " + str(exc)
    ) from exc


@dataclass(frozen=True)
class Source:
    filename: str
    category: str
    reference_bpm: float | None = None


# Solo i quattro file richiesti, ciascuno in una categoria distinta.
SOURCES = (
    Source("misure_artificial_30.xlsx", "Artificiali 30 bpm", 30.0),
    Source("misure_artificial_150.xlsx", "Artificiali 150 bpm", 150.0),
    Source("misure_pietro.xlsx", "Pietro"),
    Source("misure_manfred.xlsx", "Manfred"),
)
CATEGORIES = ("Artificiali 30 bpm", "Artificiali 150 bpm", "Pietro", "Manfred")
PAIR_STATUS = "ABBINATO_ENTRO_TOLLERANZA"
THRESHOLD_BPM = 5.0
TARGET_FRACTION = 0.95
SOFTWARE_LIMIT = 200.0
METRIC_KEYS = (
    "categoria", "righe", "coppie_valide", "escluse", "bias_bpm", "sd_bpm",
    "varianza_bpm2", "mae_bpm", "rmse_bpm", "entro_5_bpm", "frazione_entro_5",
    "target_raggiunto", "banda_inferiore_bpm", "banda_superiore_bpm",
    "percentile_2_5_bpm", "percentile_97_5_bpm", "shapiro_p_esplorativo",
    "ic95_bias_inferiore", "ic95_bias_superiore", "motivo_ic_non_disponibile",
)


def number(value):
    """None e testo non diventano zero. I valori non finiti sono non validi."""
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    return float(value) if math.isfinite(value) else None


def iso(value):
    return value.isoformat(timespec="milliseconds") if isinstance(value, datetime) else value


def read_source(source: Source, folder: Path):
    path = folder / source.filename
    if not path.is_file():
        raise ValueError(f"File richiesto non trovato: {path}")
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb["Misure"]
        iterator = ws.iter_rows(values_only=True)
        headers = next(iterator)
        required = {"N", "Istante previsto UTC", "Microbit ricevuto UTC", "Microbit BPM", "Stato"}
        if source.reference_bpm is None:
            required |= {"Garmin BPM", "Garmin allineato UTC"}
        if not required.issubset(headers):
            raise ValueError(f"Colonne mancanti in {path.name}: {required - set(headers)}")
        settings = dict(list(wb["Impostazioni"].values)[1:]) if "Impostazioni" in wb.sheetnames else {}
        tolerance = float(settings.get("tolleranza_abbinamento_secondi", 1))
        grid_tolerance = float(settings.get("tolleranza_griglia_secondi", 1))
        if not 0 <= tolerance < 5 or not 0 <= grid_tolerance < 5:
            raise ValueError(f"Tolleranza non valida in {path.name}")
        records, times, used_micro, used_ref = [], [], set(), set()
        first_time = None
        for excel_row, values in enumerate(iterator, start=2):
            if all(v is None for v in values):
                continue
            row = dict(zip(headers, values))
            timestamp = row.get("Istante previsto UTC")
            micro_time = row.get("Microbit ricevuto UTC")
            ref_time = row.get("Garmin allineato UTC")
            if isinstance(timestamp, datetime):
                first_time = first_time or timestamp
                times.append(timestamp)
            ppg = number(row.get("Microbit BPM"))
            ref = source.reference_bpm if source.reference_bpm is not None else number(row.get("Garmin BPM"))
            reasons = []
            if row.get("Microbit BPM") is None:
                reasons.append("PPG_MANCANTE")
            elif ppg is None or not 0 < ppg <= SOFTWARE_LIMIT:
                reasons.append("PPG_NON_VALIDO")
            if source.reference_bpm is None:
                if row.get("Garmin BPM") is None:
                    reasons.append("RIFERIMENTO_MANCANTE")
                elif ref is None or ref <= 0:
                    reasons.append("RIFERIMENTO_NON_VALIDO")
                if row.get("Stato") != PAIR_STATUS:
                    reasons.append("STATO_NON_ABBINATO")
                if not all(isinstance(t, datetime) for t in (timestamp, micro_time, ref_time)):
                    reasons.append("TIMESTAMP_MANCANTE")
                elif (
                    abs((micro_time - timestamp).total_seconds()) > grid_tolerance
                    or abs((ref_time - timestamp).total_seconds()) > tolerance
                    or abs((ref_time - micro_time).total_seconds()) > tolerance
                ):
                    reasons.append("FUORI_TOLLERANZA")
                if micro_time is not None and micro_time in used_micro:
                    reasons.append("TIMESTAMP_PPG_RIUTILIZZATO")
                if ref_time is not None and ref_time in used_ref:
                    reasons.append("TIMESTAMP_GARMIN_RIUTILIZZATO")
            # Nessuna esclusione per errore grande o PPG fuori dal range 30-150.
            # Per il generatore costante il Garmin non e' un riferimento richiesto.
            valid = not reasons
            if valid:
                used_micro.add(micro_time)
                used_ref.add(ref_time)
            difference = ppg - ref if valid else None
            records.append({
                "categoria": source.category,
                "file": source.filename,
                "riga_excel": excel_row,
                "n_originale": row.get("N"),
                "istante_previsto_utc": iso(timestamp),
                "microbit_ricevuto_utc": iso(micro_time),
                "garmin_allineato_utc": iso(ref_time),
                "secondi_da_inizio": (timestamp - first_time).total_seconds()
                if isinstance(timestamp, datetime) and first_time else None,
                "ppg_bpm": ppg,
                "riferimento_bpm": ref,
                "tipo_riferimento": "generatore_nominale" if source.reference_bpm is not None else "Garmin",
                "stato_originale": row.get("Stato"),
                "coppia_valida": valid,
                "motivi_esclusione": ";".join(reasons),
                "differenza_bpm": difference,
                "errore_assoluto_bpm": abs(difference) if valid else None,
                "media_coppia_bpm": (ppg + ref) / 2 if valid else None,
                "entro_5_bpm": abs(difference) <= THRESHOLD_BPM if valid else None,
                "scarto_temporale_s": (ref_time - micro_time).total_seconds()
                if source.reference_bpm is None and isinstance(ref_time, datetime)
                and isinstance(micro_time, datetime) else None,
            })
        if not records:
            raise ValueError(f"Nessuna riga di dati in {path.name}")
        return records, {
            "file": path.name,
            "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            "categoria": source.category,
            "riferimento_nominale_bpm": source.reference_bpm,
            "righe": len(records),
            "intervalli_osservati_s": sorted({round((b - a).total_seconds(), 3) for a, b in zip(times, times[1:])}),
            "impostazioni": {str(k): iso(v) for k, v in settings.items()},
            "conteggio_stati_originali": dict(Counter(r["stato_originale"] for r in records)),
        }
    finally:
        wb.close()


def summarize(records, label):
    differences = np.array([r["differenza_bpm"] for r in records if r["coppia_valida"]], dtype=float)
    n = len(differences)
    result = {key: None for key in METRIC_KEYS}
    result.update(categoria=label, righe=len(records), coppie_valide=n, escluse=len(records) - n)
    result["motivo_ic_non_disponibile"] = (
        "Non sono documentate almeno due prove ripetute indipendenti, della stessa dimensione, "
        "per la stessa combinazione persona/PCB/condizione. 30 e 150 bpm sono condizioni diverse."
    )
    if n:
        bias = float(np.mean(differences))
        sd = float(np.std(differences, ddof=1)) if n > 1 else None
        within = int(np.count_nonzero(np.abs(differences) <= THRESHOLD_BPM))
        result.update(
            bias_bpm=bias, sd_bpm=sd,
            varianza_bpm2=sd ** 2 if sd is not None else None,
            mae_bpm=float(np.mean(np.abs(differences))),
            rmse_bpm=float(np.sqrt(np.mean(differences ** 2))),
            entro_5_bpm=within, frazione_entro_5=within / n,
            target_raggiunto=within / n >= TARGET_FRACTION,
            banda_inferiore_bpm=bias - 1.96 * sd if sd is not None else None,
            banda_superiore_bpm=bias + 1.96 * sd if sd is not None else None,
            percentile_2_5_bpm=float(np.quantile(differences, 0.025)),
            percentile_97_5_bpm=float(np.quantile(differences, 0.975)),
        )
        if 3 <= n <= 5000 and sd and sd > 0:
            result["shapiro_p_esplorativo"] = float(stats.shapiro(differences).pvalue)
    return result


def write_csv(path, rows, keys=None):
    if not rows:
        return
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=keys or list(rows[0]), delimiter=";")
        writer.writeheader()
        writer.writerows(rows)


def category_slug(label):
    return label.lower().replace(" ", "_")


def make_plots(records, summary, output: Path):
    """Quattro figure separate per esperimento, con tutti i testi in inglese."""
    label = summary["categoria"]
    artificial = label.startswith("Artificiali ")
    english_label = label.replace("Artificiali", "Artificial")
    sources = list(dict.fromkeys(r["file"] for r in records))
    experiment_dir = output / category_slug(label)
    experiment_dir.mkdir(parents=True, exist_ok=True)
    plots = []

    def new_plot():
        fig, ax = plt.subplots(figsize=(8.4, 5.6))
        fig.subplots_adjust(left=.13, right=.97, bottom=.15, top=.88)
        return fig, ax

    def save_plot(fig, ax, name, title, caption):
        ax.set_title(f"{english_label} - {title}", fontsize=13, fontweight="bold", pad=14)
        ax.grid(alpha=.18)
        ax.spines[["right", "top"]].set_visible(False)
        png, pdf = experiment_dir / f"{name}.png", experiment_dir / f"{name}.pdf"
        fig.savefig(png, dpi=200, facecolor="white", bbox_inches="tight", pad_inches=.18)
        fig.savefig(pdf, facecolor="white", bbox_inches="tight", pad_inches=.18)
        plt.close(fig)
        plots.append({"png": png, "pdf": pdf, "title": title, "caption": caption})

    colors = ["#187a83", "#a45b0c"]
    fig, ax = new_plot()
    for idx, filename in enumerate(sources):
        rows = [r for r in records if r["file"] == filename]
        x = [r["secondi_da_inizio"] for r in rows]
        y = [r["ppg_bpm"] if r["ppg_bpm"] is not None else np.nan for r in rows]
        ref = [r["riferimento_bpm"] if r["riferimento_bpm"] is not None else np.nan for r in rows]
        run = f'{int(ref[0])} bpm' if artificial else label
        ax.plot(x, y, ".-", color=colors[idx % 2], label="PPG", markersize=4)
        ax.plot(x, ref, "--", color=colors[idx % 2] if artificial else "#7d5f9e",
                label=f"Reference ({run})" if artificial else "Garmin reference")
        excluded_x = [r["secondi_da_inizio"] for r in rows if not r["coppia_valida"]]
        if excluded_x:
            ax.scatter(excluded_x, [0.035 + idx * .045] * len(excluded_x), transform=ax.get_xaxis_transform(),
                       color=colors[idx % 2], marker="x", label="Missing / excluded readings", s=35)
    ax.set(xlabel="Time since start (s)", ylabel="Heart rate (bpm)")
    all_times = [r["secondi_da_inizio"] for r in records if r["secondi_da_inizio"] is not None]
    if all_times:
        pad = max(5, (max(all_times) - min(all_times)) * .035)
        ax.set_xlim(min(all_times) - pad, max(all_times) + pad)
    ax.legend(fontsize=9, loc="best")
    save_plot(fig, ax, "time_series", "Heart rate over time",
              "Andamento temporale. Le croci in basso indicano righe mancanti o escluse.")

    valid = [r for r in records if r["coppia_valida"]]
    differences = np.array([r["differenza_bpm"] for r in valid], dtype=float)
    fig, ax = new_plot()
    ax.axhspan(-THRESHOLD_BPM, THRESHOLD_BPM, color="#d4ebe7", zorder=0, label="±5 bpm target band")
    for idx, filename in enumerate(sources):
        rows = [r for r in valid if r["file"] == filename]
        ax.scatter([r["media_coppia_bpm"] for r in rows], [r["differenza_bpm"] for r in rows],
                   s=35, alpha=.7, color=colors[idx % 2], edgecolors="white", linewidths=.4, label="Paired readings")
    if summary["bias_bpm"] is not None:
        ax.axhline(summary["bias_bpm"], color="#26394b", label="Mean difference (bias)")
    if summary["banda_inferiore_bpm"] is not None:
        for key in ("banda_inferiore_bpm", "banda_superiore_bpm"):
            ax.axhline(summary[key], color="#b74f46", linestyle="--", label="Mean ±1.96 SD" if key.endswith("inferiore_bpm") else None)
    ax.set(xlabel="Mean of PPG and reference (bpm)", ylabel="PPG - reference (bpm)")
    ax.legend(fontsize=9, loc="best")
    save_plot(fig, ax, "bland_altman", "Bland-Altman (descriptive)",
              "Bland-Altman descrittivo. Più letture possono coincidere nello stesso punto. "
              "Le linee rosse tratteggiate sono bias ±1,96 SD, da interpretare con le cautele sotto.")

    fig, ax = new_plot()
    if len(differences):
        bins = min(18, max(5, int(math.sqrt(len(differences)))))
        ax.hist(differences, bins=bins, color="#187a83", edgecolor="white")
        ax.axvline(-THRESHOLD_BPM, color="#b74f46", linestyle=":")
        ax.axvline(THRESHOLD_BPM, color="#b74f46", linestyle=":", label="±5 bpm")
        ax.legend(fontsize=9)
    else:
        ax.text(.5, .5, "No valid paired readings", transform=ax.transAxes, ha="center")
    ax.set(xlabel="PPG - reference (bpm)", ylabel="Number of readings")
    ax.yaxis.set_major_locator(matplotlib.ticker.MaxNLocator(integer=True))
    save_plot(fig, ax, "histogram", "Distribution of differences",
              "Distribuzione delle differenze PPG meno riferimento. Le linee verticali indicano ±5 bpm.")

    fig, ax = new_plot()
    if len(differences) >= 3 and np.ptp(differences) > 0:
        stats.probplot(differences, dist="norm", plot=ax)
        ax.get_lines()[0].set(color="#187a83", marker="o", markersize=4)
        ax.get_lines()[1].set(color="#b74f46")
    else:
        ax.text(.5, .5, "Q-Q plot unavailable:\nconstant differences or fewer than 3 valid pairs",
                ha="center", va="center", transform=ax.transAxes)
        ax.set_xticks([])
        ax.set_yticks([])
    ax.set(xlabel="Theoretical normal quantiles", ylabel="Ordered differences (bpm)")
    save_plot(fig, ax, "qq_plot", "Normal Q-Q plot",
              "Q-Q plot delle differenze rispetto alla distribuzione normale. Non applicabile se le differenze sono costanti.")
    return plots


def fmt(value, places=2):
    if value is None:
        return "n.d."
    return f"{value:.{places}f}".replace(".", ",")


def table(headers, rows):
    return '<div class="table-wrap"><table><thead><tr>' + "".join(f"<th>{escape(str(h))}</th>" for h in headers) + "</tr></thead><tbody>" + "".join(
        "<tr>" + "".join(f"<td>{escape(str(v))}</td>" for v in row) + "</tr>" for row in rows
    ) + "</tbody></table></div>"


def summary_row(s):
    return [s["categoria"], s["coppie_valide"], s["escluse"], fmt(s["bias_bpm"]),
            fmt(s["sd_bpm"]), fmt(s["mae_bpm"]), fmt(s["rmse_bpm"]),
            f'{s["entro_5_bpm"]}/{s["coppie_valide"]} ({fmt(100 * s["frazione_entro_5"], 1)}%)'
            if s["frazione_entro_5"] is not None else "n.d.",
            "Rispettato" if s["target_raggiunto"] else ("Non rispettato" if s["target_raggiunto"] is not None else "Non valutabile")]


def make_report(output, records, summaries, manifest, plot_paths, generated):
    dates = sorted({r["istante_previsto_utc"][:10] for r in records if r["istante_previsto_utc"]})
    period = " – ".join(dict.fromkeys((dates[0], dates[-1]))) if dates else "non documentato"
    heads = ["Gruppo", "Valide", "Escluse", "Bias (bpm)", "SD (bpm)", "MAE (bpm)", "RMSE (bpm)", "Entro ±5 bpm", "Target ≥95%"]
    content = [f"""<!doctype html><html lang="it"><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1"><title>Lab 6 - Analisi delle misure</title>
<style>
:root{{color-scheme:light}}*{{box-sizing:border-box}}body{{margin:0;background:#f3f6f7;color:#223745;font:16px/1.6 'Segoe UI',Arial,sans-serif}}
main{{max-width:1220px;margin:auto;padding:40px 30px 64px}}header{{border-top:6px solid #187a83;padding:25px 0 15px}}h1{{font-size:34px;line-height:1.2;margin:8px 0}}h2{{font-size:25px;margin:0 0 14px;color:#16394b}}h3{{font-size:19px;margin:25px 0 10px}}
.eyebrow{{text-transform:uppercase;letter-spacing:2px;color:#187a83;font-size:12px;font-weight:700}}.muted{{color:#5a6b74;font-size:14px}}
nav{{display:flex;gap:22px;flex-wrap:wrap;margin:18px 0}}a{{color:#126a74}}section{{background:white;border:1px solid #dce5e9;border-radius:9px;padding:26px;margin-top:24px;scroll-margin-top:20px}}
.note{{border-left:4px solid #c48b3a;padding:10px 16px;background:#fff8eb;margin:18px 0}}.table-wrap{{overflow-x:auto}}table{{border-collapse:collapse;width:100%;font-size:14px;margin:8px 0 16px}}th{{text-align:left;background:#eaf1f3;color:#16394b;white-space:normal}}td,th{{padding:10px 11px;border-bottom:1px solid #dce5e9;vertical-align:top}}td:not(:first-child){{font-variant-numeric:tabular-nums}}figure{{margin:22px 0}}figure img{{width:100%;height:auto}}figcaption{{font-size:13px;color:#5a6b74}}details{{margin-top:18px}}summary{{cursor:pointer;font-weight:600}}code{{background:#edf2f4;padding:2px 5px;font-size:13px}}li{{margin:6px 0}}.metrics{{display:grid;grid-template-columns:repeat(3,1fr);gap:14px;margin:20px 0}}.metric{{background:#eaf3f3;border-radius:6px;padding:15px}}.metric strong{{display:block;font-size:26px;color:#126a74}}.metric span{{font-size:13px}}
.plots{{display:grid;grid-template-columns:repeat(2,minmax(0,1fr));gap:22px}}.plots figure{{margin:0}}.plots figcaption a{{white-space:nowrap}}
@media(max-width:700px){{main{{padding:20px 12px}}section{{padding:18px}}h1{{font-size:28px}}.metrics,.plots{{grid-template-columns:1fr}}}}
@media print{{body{{background:white;font-size:10pt}}main{{max-width:none;padding:0}}section{{break-inside:avoid;border-radius:0;margin:12px 0;padding:15px}}nav{{display:none}}details{{display:block}}figure{{break-inside:avoid}}th,td{{padding:6px}}}}
</style><main><header><div class="eyebrow">KTH · Lab 6</div><h1>Analisi delle misure di frequenza cardiaca</h1>
<p>Quattro categorie separate: artificiali a 30 bpm, artificiali a 150 bpm, Pietro e Manfred.</p>
<p class="muted">Date dei dati (UTC): {escape(period)}. Report generato il {escape(generated)}. Unità delle differenze: bpm.</p>
<nav><a href="#risultati">Risultati</a><a href="#artificiali_30_bpm">Artificiali 30 bpm</a><a href="#artificiali_150_bpm">Artificiali 150 bpm</a><a href="#pietro">Pietro</a><a href="#manfred">Manfred</a><a href="#metodo">Metodo e limiti</a></nav></header>
<section id="risultati"><h2>Risultati per categoria</h2>
<p>La differenza è <strong>PPG − riferimento</strong>. Un bias positivo indica letture PPG mediamente più alte.
Per le persone il riferimento è Garmin. Per gli artificiali si usano i valori nominali del generatore.</p>"""]
    content.append(table(heads, [summary_row(s) for s in summaries]))
    content.append("<p class='muted'>SD campionaria con denominatore n−1. Le statistiche a 30 e 150 bpm sono calcolate separatamente.</p>")
    content.append("<p>" + " ".join(
        f"<strong>{escape(s['categoria'])}:</strong> {fmt(100 * s['frazione_entro_5'], 1)}% entro ±5 bpm."
        for s in summaries
    ) + " L'esito riguarda le letture disponibili e le condizioni provate.</p>")
    content.append("<div class='note'><strong>Interpretazione.</strong> Il confronto umano misura l'accordo con Garmin, non l'accuratezza assoluta. "
                   "Gli artificiali usano 30 e 150 bpm desunti dai nomi dei file: la verifica indipendente del generatore non è documentata. "
                   "L'esito artificiale resta condizionato a questa verifica.</div></section>")

    for s, experiment_plots in zip(summaries, plot_paths):
        category = s["categoria"]
        rows = [r for r in records if r["categoria"] == category]
        content.append(f'<section id="{category_slug(category)}"><h2>{category}</h2>')
        content.append('<div class="metrics">'
                       f'<div class="metric"><strong>{s["coppie_valide"]}/{s["righe"]}</strong><span>Coppie valide / righe presenti</span></div>'
                       f'<div class="metric"><strong>{fmt(s["bias_bpm"])} bpm</strong><span>Bias PPG − riferimento</span></div>'
                       f'<div class="metric"><strong>{fmt(100 * s["frazione_entro_5"], 1) if s["frazione_entro_5"] is not None else "n.d."}%</strong><span>Letture valide entro ±5 bpm</span></div></div>')
        content.append('<div class="plots">')
        for plot in experiment_plots:
            encoded = base64.b64encode(plot["png"].read_bytes()).decode("ascii")
            png_link = plot["png"].relative_to(output).as_posix()
            pdf_link = plot["pdf"].relative_to(output).as_posix()
            content.append(f'<figure><img alt="{escape(category)}: {escape(plot["title"])}" src="data:image/png;base64,{encoded}">'
                           f'<figcaption>{escape(plot["caption"])} '
                           f'<a href="{png_link}">PNG</a> · <a href="{pdf_link}">PDF</a></figcaption></figure>')
        content.append('</div>')
        content.append('<h3>Distribuzione e limiti descrittivi</h3>')
        p = s["shapiro_p_esplorativo"]
        p_text = "non calcolabile: differenze costanti o meno di 3 coppie" if p is None else ("<0,0001" if p < .0001 else fmt(p, 4))
        content.append(table(["Indicatore", "Valore"], [
            ["Varianza delle differenze (bpm²)", fmt(s["varianza_bpm2"])],
            ["Bande aritmetiche bias ±1,96 SD (bpm)", f'{fmt(s["banda_inferiore_bpm"])} ; {fmt(s["banda_superiore_bpm"])}'],
            ["Percentili empirici 2,5 e 97,5 delle differenze (bpm)", f'{fmt(s["percentile_2_5_bpm"])} ; {fmt(s["percentile_97_5_bpm"])}'],
            ["p di Shapiro-Wilk (esplorativo)", p_text],
            ["IC 95% del bias da medie di prove ripetute", "Non stimabile dai file disponibili"],
        ]))
        if s["sd_bpm"] == 0:
            content.append("<p>Le differenze osservate sono costanti: la SD è zero e il Q-Q plot e il test di normalità non sono applicabili. "
                           "Le bande coincidono con il bias e descrivono soltanto questa prova; non dimostrano una copertura del 95% in nuove misure.</p>")
        elif p is not None and p < .05:
            content.append("<p>Il Q-Q plot e il test esplorativo segnalano una distribuzione non compatibile con una semplice approssimazione normale. "
                           "Le bande bias ±1,96 SD sono mostrate come calcolo descrittivo e non come limiti di accordo con copertura del 95% dimostrata. "
                           "La stabilità della varianza non è stata verificata in prove indipendenti.</p>")
        else:
            content.append("<p>La copertura del 95% delle bande bias ±1,96 SD richiede differenze approssimativamente normali e varianza stabile. "
                           "Il grafico e il test esplorativo non dimostrano tali condizioni; la stabilità della varianza resta da verificare.</p>")
        content.append("<p class='muted'>Shapiro-Wilk è solo un ausilio descrittivo: la dipendenza temporale delle letture può invalidarne l'interpretazione inferenziale. "
                       "I percentili sono descrittivi del campione e non intervalli di confidenza.</p>")
        content.append('<h3>Letture mancanti e non valide</h3>')
        excluded = [r for r in rows if not r["coppia_valida"]]
        exclusion_text = "1 riga esclusa" if len(excluded) == 1 else f"{len(excluded)} righe escluse"
        content.append(f'<p>{exclusion_text} su {len(rows)} presenti ({fmt(100 * len(excluded) / len(rows), 1)}%). '
                       'Le percentuali entro ±5 bpm usano solo le coppie valide. Nessuna imputazione e nessuna rimozione degli errori grandi.</p>')
        if excluded:
            content.append(table(["File", "N originale", "Motivi di esclusione"], [
                [r["file"], r["n_originale"], r["motivi_esclusione"]] for r in excluded
            ]))
        content.append('<details><summary>Mostra tutte le differenze e gli stati</summary>')
        content.append(table(["File", "N", "Tempo (s)", "PPG", "Riferimento", "d (bpm)", "Valida", "Stato originale"], [
            [r["file"], r["n_originale"], fmt(r["secondi_da_inizio"], 0), fmt(r["ppg_bpm"], 0),
             fmt(r["riferimento_bpm"], 0), fmt(r["differenza_bpm"]), "Sì" if r["coppia_valida"] else "No", r["stato_originale"]]
            for r in rows
        ]))
        content.append('</details></section>')

    source_rows = []
    for m in manifest:
        valid_count = sum(r["coppia_valida"] for r in records if r["file"] == m["file"])
        source_rows.append([m["file"], m["righe"], valid_count, ", ".join(str(x) for x in m["intervalli_osservati_s"]),
                            m["impostazioni"].get("stato_registrazione", "Non documentato")])
    content.append('<section id="metodo"><h2>Metodo, protocollo e limiti</h2>')
    content.append(table(["Fonte", "Righe", "Valide", "Passo osservato (s)", "Registrazione"], source_rows))
    by_category = {s["categoria"]: s for s in summaries}
    pietro, manfred = by_category["Pietro"], by_category["Manfred"]
    condition_text = " ".join(
        f"{escape(m['file'])}: {m['righe']} righe, {sum(r['coppia_valida'] for r in records if r['file'] == m['file'])} valide, "
        f"stato '{escape(str(m['impostazioni'].get('stato_registrazione', 'non documentato')))}'."
        for m in manifest if m['riferimento_nominale_bpm'] is not None
    )
    interval_text = "; ".join(f"{m['file']}: {m['intervalli_osservati_s']} s" for m in manifest)
    plan_counts = (
        f"<li>Il piano specifica un passo di 5 s. Passi osservati: {escape(interval_text)}. Nessun ricampionamento o interpolazione.</li>"
        "<li>Il PDF prevede 100 osservazioni artificiali e non ne specifica la ripartizione tra i livelli. "
        "I conteggi e le statistiche sono riportati separatamente per 30 e 150 bpm.</li>"
        f"<li>{condition_text}</li>"
        f"<li>Persone: Pietro {pietro['righe']} righe e {pietro['coppie_valide']} coppie valide; Manfred {manfred['righe']} righe e "
        f"{manfred['coppie_valide']} coppie valide. Il piano richiede 40 coppie valide per combinazione PCB/persona.</li>"
    )
    content.append("""<h3>Calcoli e selezione delle coppie</h3>
<p>dᵢ = PPGᵢ − riferimentoᵢ; bias = media(d); SD = √[Σ(dᵢ−bias)²/(n−1)]; MAE = media(|d|); RMSE = √media(d²).
La quota entro ±5 bpm include esattamente −5 e +5. Il target è quota ≥0,95 sui validi e non è un requisito normativo.</p>
<p>Per le persone si richiedono BPM numerici positivi, PPG non oltre il limite software di 200 bpm, stato ABBINATO_ENTRO_TOLLERANZA,
timestamp disponibili e scarti entro le tolleranze registrate (1 s in queste fonti). Non si riutilizzano timestamp all'interno dello stesso file.
Non si stimano ritardi dai BPM. Per gli artificiali il riferimento è costante e non richiede Garmin: lo stato originale GARMIN_DA_SCARICARE
non invalida una lettura PPG confrontabile con il generatore. I valori PPG oltre 150 bpm sono conservati per non selezionare soltanto gli errori piccoli.</p>
<h3>Intervallo di confidenza del bias</h3>
<p>Il piano richiede un IC basato su medie di prove ripetute indipendenti di uguale dimensione:
media delle medie ± t<sub>0,975;k−1</sub> × SD(medie)/√k.
È disponibile un solo file per persona, senza identificazione del PCB e senza ripetizioni indipendenti documentate.
Le condizioni artificiali a 30 e 150 bpm hanno dimensioni diverse e non sono ripetizioni della stessa condizione.
L'IC resta non disponibile; i campioni consecutivi non vengono trattati come soggetti indipendenti e non vengono divisi in finte repliche.</p>
<h3>Differenze rispetto al piano</h3><ul>""" + plan_counts + """
<li>Il piano prevede due PCB su entrambe le persone. I file non distinguono il PCB e l'indirizzo Bluetooth del micro:bit non identifica la scheda analogica. Non è possibile confrontare i due PCB o verificare la completezza della campagna.</li>
<li>Sono provati solo i due estremi artificiali, non l'intero intervallo 30–150 bpm. Il limite software a 200 bpm non estende l'intervallo validato.</li>
<li>Modello dichiarato nel README del progetto: Garmin vívoactive 3 Music. La sorgente cardiaca effettiva (polso o fascia), l'averaging/ritardo Garmin e la verifica del generatore non sono documentati nei quattro file.</li>
<li>Il piano descrive adulti sani, seduti e fermi dopo 5 minuti di riposo, luce e temperatura ambiente. I file non consentono di verificare queste condizioni o il riposizionamento tra prove.</li>
<li>Il micro:bit usa il timestamp di ricezione sul PC, Garmin quello del record FIT: abbinamento temporale approssimato, senza certificare simultaneità fisica. I risultati valgono per le persone, i dispositivi e le condizioni effettivamente provati.</li>
</ul><h3>Fonti e riproducibilità</h3>
<p>Protocollo: <em>Lab6 - Plan and report.pdf</em>, pagine 1–2, fornito con la richiesta. Dati: i quattro file elencati sopra, fogli Misure e Impostazioni.
README.md è usato soltanto per il modello Garmin dichiarato. Gli originali non vengono modificati.
manifest.json conserva hash SHA-256, impostazioni e versioni delle librerie; dati_analizzati.csv conserva gli stati originali e i motivi di esclusione.
riepilogo_categorie.csv e dettaglio_artificiali.csv riportano i valori numerici non arrotondati.</p>
<p class="muted">Il report HTML contiene le immagini e si può leggere offline. Ciascun esperimento ha una propria cartella con quattro grafici separati,
ognuno in PNG e PDF vettoriale. Titoli, assi, legende e annotazioni delle immagini sono in inglese.</p>
</section></main></html>""")
    (output / "report_analisi_lab6.html").write_text("\n".join(content), encoding="utf-8")


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input-dir", type=Path, default=ROOT, help="Cartella dei quattro Excel originali")
    parser.add_argument("--output-dir", type=Path, default=ROOT / "risultati_analisi", help="Cartella dei risultati")
    args = parser.parse_args(argv)
    records, manifest = [], []
    for source in SOURCES:
        data, info = read_source(source, args.input_dir)
        records.extend(data)
        manifest.append(info)
    summaries = [summarize([r for r in records if r["categoria"] == category], category) for category in CATEGORIES]
    endpoints = [s for s in summaries if s["categoria"].startswith("Artificiali ")]
    if any(s["coppie_valide"] == 0 for s in summaries):
        raise ValueError("Una categoria non ha coppie valide: controllare fonti e motivi di esclusione.")
    output = args.output_dir.resolve()
    source_paths = {(args.input_dir / s.filename).resolve() for s in SOURCES}
    if output in source_paths:
        raise ValueError("La destinazione non può essere un file sorgente.")
    output.mkdir(parents=True, exist_ok=True)
    generated = datetime.now(timezone.utc).isoformat(timespec="seconds")
    versions = {name: importlib.metadata.version(name) for name in ["openpyxl", "numpy", "scipy", "matplotlib"]}
    plot_paths = [make_plots([r for r in records if r["categoria"] == s["categoria"]], s, output) for s in summaries]
    write_csv(output / "dati_analizzati.csv", records)
    write_csv(output / "riepilogo_categorie.csv", summaries, METRIC_KEYS)
    write_csv(output / "dettaglio_artificiali.csv", endpoints, METRIC_KEYS)
    result = {"categorie": summaries, "dettaglio_artificiali": endpoints}
    (output / "risultati.json").write_text(json.dumps(result, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8")
    (output / "manifest.json").write_text(json.dumps({"generato_utc": generated, "python": sys.version,
        "librerie": versions, "fonti": manifest}, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8")
    make_report(output, records, summaries, manifest, plot_paths, generated)
    for s in summaries:
        print(f'{s["categoria"]}: {s["coppie_valide"]}/{s["righe"]} valide, bias {s["bias_bpm"]:.3f} bpm, '
              f'MAE {s["mae_bpm"]:.3f} bpm, entro +/-5: {s["frazione_entro_5"]:.1%}')
    print(f'Report: {output / "report_analisi_lab6.html"}')
    return result


if __name__ == "__main__":
    main()
